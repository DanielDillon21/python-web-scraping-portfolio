"""
Turns the raw crawl (output/products.csv) into a client-ready Excel report.

Sheets:
    Summary           headline numbers
    Products          full cleaned dataset, filterable, frozen header
    By Category       product count, average price, average rating, total stock
    Low Stock         items with 3 or fewer units left
    Top Rated Value   5-star books sorted by price

Run:
    python build_report.py
Output:
    output/catalog_report.xlsx
"""

from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path(__file__).parent / "output"
SOURCE = OUTPUT_DIR / "products.csv"
REPORT = OUTPUT_DIR / "catalog_report.xlsx"

LOW_STOCK_UNITS = 3

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def load():
    df = pd.read_csv(SOURCE)
    df = df.drop_duplicates(subset="upc").sort_values(["category", "title"])
    return df.reset_index(drop=True)


def build_tables(df):
    by_category = (
        df.groupby("category")
        .agg(
            products=("upc", "count"),
            avg_price_gbp=("price_gbp", "mean"),
            avg_rating=("rating", "mean"),
            total_stock=("in_stock", "sum"),
        )
        .round(2)
        .sort_values("products", ascending=False)
        .reset_index()
    )

    low_stock = df[df["in_stock"] <= LOW_STOCK_UNITS].sort_values("in_stock")[
        ["title", "category", "price_gbp", "in_stock", "url"]
    ]

    top_rated_value = df[df["rating"] == 5].sort_values("price_gbp")[
        ["title", "category", "price_gbp", "in_stock", "url"]
    ]

    summary = pd.DataFrame(
        {
            "metric": [
                "Products scraped",
                "Categories",
                "Average price (GBP)",
                "Cheapest (GBP)",
                "Most expensive (GBP)",
                "Average star rating",
                f"Low-stock items (<= {LOW_STOCK_UNITS} units)",
                "5-star items",
            ],
            "value": [
                len(df),
                df["category"].nunique(),
                round(df["price_gbp"].mean(), 2),
                df["price_gbp"].min(),
                df["price_gbp"].max(),
                round(df["rating"].mean(), 2),
                len(low_stock),
                len(top_rated_value),
            ],
        },
        dtype=object,  # keep counts as integers instead of upcasting to 1000.00
    )

    return {
        "Summary": summary,
        "Products": df.drop(columns=["description"]),
        "By Category": by_category,
        "Low Stock": low_stock,
        "Top Rated Value": top_rated_value,
    }


def style_sheet(ws, df):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 2, 60)


def add_category_chart(ws, rows):
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Top 15 categories by product count"
    chart.y_axis.title = "Products"
    chart.height, chart.width = 12, 22
    top = min(rows, 15)
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=top + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=top + 1))
    chart.legend = None
    ws.add_chart(chart, "G2")


def main():
    df = load()
    tables = build_tables(df)
    with pd.ExcelWriter(REPORT, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)
            style_sheet(writer.sheets[name], table)
        add_category_chart(writer.sheets["By Category"], len(tables["By Category"]))
    print(f"Wrote {REPORT} ({len(df)} products, {df['category'].nunique()} categories)")


if __name__ == "__main__":
    main()
